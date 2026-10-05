"""Write what Xero returned as a workbook the rest of fma already knows how to read.

The shape is the one a Xero export has, because `fma read-ledger` is the gate every
pack already runs on its inputs:

    row 1   the report's title, as Xero wrote it
    row 2   the organisation, as Xero wrote it
    row 3   ONE date line: "As at 30 June 2026" or "1 July 2025 to 30 June 2026"
    row 4   (blank)
    row 5   the column headers, first cell "Account"
    ...     sections, lines and totals, verbatim, in Xero's order

Differences from a hand export, each deliberate:
  * Totals are the values Xero returned, not formulas. Xero's hand exports write every
    total as a formula with a cached zero; there is nothing to gain by recreating the
    trap read-ledger exists to defuse. The workbook says who wrote it (`creator`), so
    read-ledger can tell a pull from a file somebody re-saved in Excel.
  * One report per workbook, one date line per report. read-ledger's date gate reads
    every sheet's header, and a second date-bearing line would be a second claim.
  * Nothing is added that Xero did not say: no ageing columns, no reclassification.

A file is written under a temporary name, read back through read-ledger's own loader,
checked against what was meant, hashed, and only then moved into place. A synced folder
never sees half a file, and a file this tool cannot read back is never kept.
"""

from __future__ import annotations

import contextlib
import hashlib
import os
import re
from decimal import Decimal
from pathlib import Path

from .. import __version__
from ..errors import InputProblem, Refusal
from ..read_ledger import loader, metadata
from .reports import Report

CREATOR = f"fma xero {__version__}"
_MONEY_FORMAT = "#,##0.00;-#,##0.00"


def grid_for(report: Report, date_line: str) -> list[list]:
    header = list(report.header) or [""]
    if not header[0].strip():
        header[0] = "Account"
    rows: list[list] = [[report.name], [report.entity], [date_line], [], header, []]
    for sec in report.sections:
        if sec.title:
            rows.append([sec.title])
        for ln in sec.lines:
            rows.append([ln.label, *ln.values])
        rows.append([])
    while rows and not rows[-1]:
        rows.pop()
    return rows


def table_grid(title: str, entity: str, note: str, header: list[str],
               body: list[list]) -> list[list]:
    """A plain listing (a chart of accounts) under the same three-line head. `note`
    must not read as a date line: a listing is as at the moment it was pulled."""
    return [[title], [entity], [note], [], list(header), *[list(r) for r in body]]


def name_reads_as_itself(entity: str) -> bool:
    """Whether read-ledger's header parser would take this organisation name for a
    name. A name that is itself a date line ("Old File to 30 June 2023", "As at ...")
    is read as the report's date instead, and every workbook for that organisation
    would fail its read-back. Asked before a pull spends a single report call."""
    probe = [["Report"], [entity], ["As at 1 January 2000"], [], ["Account", "x"]]
    meta = metadata.extract(probe)
    return (meta.get("entity") == str(entity).strip()
            and meta.get("report_date") == "2000-01-01"
            and not meta.get("report_period"))


def _sheet_title(name: str) -> str:
    return (re.sub(r"[\[\]:*?/\\]", " ", name).strip() or "Report")[:31]


def write_workbook(path: Path, sheets: list[tuple[str, list[list]]],
                   expect: dict | None = None) -> dict:
    """Write `sheets` ([(title, grid)]) to `path` atomically; return what was written.

    `expect` holds what the first sheet's header must read back as through
    read-ledger's own parser: {"date": iso} for an as-at report, {"start": iso,
    "end": iso} for a period, or None for a listing that claims no date."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    path = Path(path)
    if not path.is_absolute():
        raise InputProblem("PATH_NOT_ABSOLUTE", f"{path} is not an absolute path")
    tmp = path.with_name(f".{path.stem}.tmp.xlsx")      # the loader opens .xlsx only
    wb = openpyxl.Workbook()
    wb.properties.creator = CREATOR
    wb.remove(wb.active)
    used = set()
    for title, grid in sheets:
        name, n = _sheet_title(title), 2
        while name in used:
            name, n = f"{_sheet_title(title)[:28]} {n}", n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        widths: dict[int, int] = {}
        for r, row in enumerate(grid, start=1):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, Decimal):
                    cell.number_format = _MONEY_FORMAT
                elif isinstance(value, str) and value.startswith("="):
                    # a label that begins with "=" is text; openpyxl would make it
                    # a formula and read-ledger would then try to evaluate it
                    cell.data_type = "s"
                widths[c] = max(widths.get(c, 0), len(str(value)))
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 12), 60)
    try:
        with contextlib.suppress(FileNotFoundError):
            tmp.unlink()
        try:
            wb.save(tmp)
        except OSError as e:
            raise InputProblem("CANNOT_WRITE",
                               f"cannot write into {path.parent} ({e.strerror or e}) -- "
                               "if the folder lives in OneDrive it may be cloud-only; "
                               "open it in Finder first")
        grids = loader.load(tmp)
        meta = metadata.extract(grids[0].grid)
        _check_read_back(meta, sheets[0][1], expect, path.name)
        digest = hashlib.sha256(tmp.read_bytes()).hexdigest()
        os.replace(tmp, path)
    except BaseException:
        with contextlib.suppress(FileNotFoundError):
            tmp.unlink()
        raise
    return {"file": path.name, "sha256": digest,
            "rows": sum(1 for row in sheets[0][1] if any(v is not None for v in row)),
            "report_date": meta.get("report_date"),
            "report_period": meta.get("report_period")}


def _check_read_back(meta: dict, grid: list[list], expect: dict | None, name: str) -> None:
    def refuse(why: str):
        raise Refusal("READ_BACK_MISMATCH",
                      f"{name} did not read back as written ({why}); nothing was kept")
    # read-ledger strips the text of a header cell, so the comparison does too
    title = str(grid[0][0]).strip() if grid and grid[0] else None
    entity = str(grid[1][0]).strip() if len(grid) > 1 and grid[1] else None
    if meta.get("report_title") != title:
        refuse(f"title {meta.get('report_title')!r}, wrote {title!r}")
    if entity and meta.get("entity") != entity:
        refuse(f"organisation {meta.get('entity')!r}, wrote {entity!r}")
    period = meta.get("report_period") or {}
    if expect is None:
        if meta.get("report_date") or period.get("end"):
            refuse("a listing read back as if it carried a report date")
    elif "date" in expect:
        if meta.get("report_date") != expect["date"]:
            refuse(f"as-at {meta.get('report_date')!r}, wrote {expect['date']!r}")
    else:
        if (period.get("start"), period.get("end")) != (expect["start"], expect["end"]):
            refuse(f"period {period!r}, wrote {expect['start']} to {expect['end']}")


def write_bytes(path: Path, blob: bytes) -> str:
    """A raw response, kept exactly as Xero sent it. Returns its sha256."""
    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp")
    try:
        tmp.write_bytes(blob)
        os.replace(tmp, path)
    except OSError as e:
        with contextlib.suppress(FileNotFoundError):
            tmp.unlink()
        raise InputProblem("CANNOT_WRITE", f"cannot write {path.name} into {path.parent} "
                                           f"({e.strerror or e})")
    return hashlib.sha256(blob).hexdigest()
