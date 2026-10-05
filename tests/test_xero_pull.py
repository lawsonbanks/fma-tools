"""xero pull: a typed date or nothing, proved dates, ties before files, all or none."""

import hashlib
import json
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from fma_tools.xero import reports
from xero_fakes import Ledger, Org, signed_in, three_orgs

AS_AT = "2026-06-30"


@pytest.fixture
def two(xero):
    """Two invented companies, signed in, keyed A and B."""
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    xero.add_org(Org("t-b", "Entity B Pty Ltd", ledger=Ledger(
        bank=[("090", "Business Bank Account", 2000)],
        current_assets=[("610", "Accounts Receivable", 4000)],
        current_liabilities=[("800", "Accounts Payable", 1000)],
        income=[("200", "Sales", 9000)],
        expenses=[("400", "Advertising", 2000), ("405", "Bank Charges", 500)])))
    signed_in(xero, ["t-a", "t-b"], keys={"t-a": "A", "t-b": "B"})
    return xero


def _pull(run_cli, out, *extra):
    return run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--prefix", "ACME",
                    "--out", str(out), *extra])


def _files(out):
    return sorted(p.name for p in out.iterdir() if p.is_file())


# -- the date is typed, or there is no pull ------------------------------------------

def test_no_as_at_is_a_refusal_that_says_why(two, run_cli, tmp_path):
    code, env = run_cli(["xero", "pull", "--all", "--out", str(tmp_path / "p")])
    assert code == 1 and env["status"] == "refuse"
    assert env["problems"][0]["code"] == "AS_AT_REQUIRED"
    assert "never chooses a date" in env["problems"][0]["message"]
    assert not two.api_calls(), "nothing may be asked of Xero without a date"


def test_a_bad_date_is_exit_2(two, run_cli, tmp_path):
    code, env = run_cli(["xero", "pull", "--as-at", "30/06/2026", "--all",
                         "--out", str(tmp_path / "p")])
    assert code == 2 and env["problems"][0]["code"] == "DATE_INVALID"


def test_naming_no_organisation_is_a_refusal_even_with_one_connected(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--out", str(tmp_path / "p")])
    assert code == 1 and env["problems"][0]["code"] == "ORG_REQUIRED"
    assert "Entity A Pty Ltd" in env["problems"][0]["message"]


def test_part_of_a_name_is_never_enough(two, run_cli, tmp_path):
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "Entity",
                         "--out", str(tmp_path / "p")])
    assert code == 1 and env["problems"][0]["code"] == "ORG_UNKNOWN"
    assert "Did you mean A (Entity A Pty Ltd), B (Entity B Pty Ltd)?" in \
        env["problems"][0]["message"]
    assert not two.api_calls()


def test_a_key_that_matches_nothing_never_lands_on_the_one_company_connected(xero, run_cli, tmp_path):
    # A stale or mistyped key ("A") sits inside the only connected name. It must not
    # be taken for it: that pulled, and once disconnected, the wrong company's books.
    xero.add_org(Org("t-b", "Bravo Trading Pty Ltd"))
    signed_in(xero, ["t-b"], keys={"t-b": "BT"})
    out = tmp_path / "p"
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "A", "--reports", "tb",
                         "--out", str(out)])
    assert code == 1 and env["problems"][0]["code"] == "ORG_UNKNOWN"
    assert not out.exists() and not xero.api_calls()
    code, env = run_cli(["xero", "disconnect", "--org", "a"])
    assert code == 1 and env["problems"][0]["code"] == "ORG_UNKNOWN"
    assert len(xero.users["user-1"]["connections"]) == 1, "still connected at Xero"
    # the whole name, any case, and the key both still work
    for name in ("bravo trading pty ltd", "bt"):
        code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", name, "--reports",
                             "tb", "--out", str(tmp_path / name.replace(" ", "_"))])
        assert code == 0, env["problems"]


def test_two_companies_with_one_name_must_be_told_apart_by_key(xero, run_cli, tmp_path):
    xero.add_org(Org("t-1", "Demo Company (AU)"))
    xero.add_org(Org("t-2", "Demo Company (AU)"))
    signed_in(xero, ["t-1", "t-2"], keys={"t-1": "ONE", "t-2": "TWO"})
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "Demo Company (AU)",
                         "--out", str(tmp_path / "p")])
    assert code == 1 and env["problems"][0]["code"] == "ORG_AMBIGUOUS"
    assert "ONE (Demo Company (AU))" in env["problems"][0]["message"]


def test_keys_that_make_the_same_file_name_refuse_before_anything_is_fetched(two, run_cli, tmp_path):
    from fma_tools.xero import tenants
    rows = tenants.registry()            # as an older version, or a hand, might leave it
    rows["t-a"]["key"], rows["t-b"]["key"] = "NSW", "NSW_"
    tenants.save_registry(rows)
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "ORG_KEY_CLASH"
    assert not out.exists() and not two.api_calls()


def test_a_relative_out_is_exit_2(two, run_cli):
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", "pull-here"])
    assert code == 2 and env["problems"][0]["code"] == "PATH_NOT_ABSOLUTE"


def test_a_folder_that_already_holds_files_refuses(two, run_cli, tmp_path):
    out = tmp_path / "p"
    out.mkdir()
    (out / "earlier.txt").write_text("an earlier pull")
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "OUT_NOT_EMPTY"
    assert _files(out) == ["earlier.txt"]


# -- what a good pull leaves behind --------------------------------------------------

def test_every_dated_file_passes_read_ledger_at_its_own_date(two, run_cli, tmp_path):
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 0, env["problems"]
    record = json.loads((out / "PULL.json").read_text())
    checked = 0
    for org in record["organisations"]:
        for f in org["files"]:
            expect = f.get("as_at") or f.get("end")
            if not expect:
                continue
            code, env = run_cli(["read-ledger", str(out / f["file"]), "--expect-date", expect])
            assert code == 0, (f["file"], env["problems"])
            assert env["data"]["metadata"]["written_by"].startswith("fma xero")
            assert not any("no live formulas" in w for w in env["warnings"]), f["file"]
            code, env = run_cli(["read-ledger", str(out / f["file"]),
                                 "--expect-date", "2026-07-31"])
            assert code == 1 and env["problems"][0]["code"] == "DATE_MISMATCH", f["file"]
            checked += 1
    assert checked == 2 * 6          # per company: TB, BS, three P&Ls, bank summary


def test_names_follow_the_existing_pulls(two, run_cli, tmp_path):
    out = tmp_path / "p"
    code, _ = _pull(run_cli, out)
    assert code == 0
    assert _files(out) == sorted([
        "PULL.json", "PULL.md",
        "ACME_A_Trial_Balance_as_at_2026-06-30.xlsx",
        "ACME_A_Balance_Sheet_as_at_2026-06-30.xlsx",
        "ACME_A_Profit_and_Loss_2026-06-01_to_2026-06-30.xlsx",
        "ACME_A_Profit_and_Loss_2025-07-01_to_2026-06-30.xlsx",
        "ACME_A_Profit_and_Loss_2024-07-01_to_2025-06-30.xlsx",
        "ACME_A_Bank_Summary_2026-06-01_to_2026-06-30.xlsx",
        "ACME_A_Chart_of_Accounts.xlsx",
        "ACME_B_Trial_Balance_as_at_2026-06-30.xlsx",
        "ACME_B_Balance_Sheet_as_at_2026-06-30.xlsx",
        "ACME_B_Profit_and_Loss_2026-06-01_to_2026-06-30.xlsx",
        "ACME_B_Profit_and_Loss_2025-07-01_to_2026-06-30.xlsx",
        "ACME_B_Profit_and_Loss_2024-07-01_to_2025-06-30.xlsx",
        "ACME_B_Bank_Summary_2026-06-01_to_2026-06-30.xlsx",
        "ACME_B_Chart_of_Accounts.xlsx",
    ])
    assert not list(out.glob(".*tmp*")), "no temporary file may be left behind"


def test_the_workbook_is_the_shape_a_xero_export_has(two, run_cli, tmp_path):
    out = tmp_path / "p"
    _pull(run_cli, out)
    ws = openpyxl.load_workbook(out / "ACME_A_Balance_Sheet_as_at_2026-06-30.xlsx").active
    col_a = [c.value for c in ws["A"]]
    assert col_a[:5] == ["Balance Sheet", "Entity A Pty Ltd", "As at 30 June 2026", None,
                         "Account"]
    assert "Current Year Earnings" in col_a          # the month-end control total
    assert "Net Assets" in col_a and "Total Equity" in col_a
    row = {c.value: i for i, c in enumerate(ws["A"], start=1)}
    assert ws.cell(row=row["Net Assets"], column=2).value == 14000
    assert ws.cell(row=row["Current Year Earnings"], column=2).value == 12000
    pl = openpyxl.load_workbook(
        out / "ACME_A_Profit_and_Loss_2025-07-01_to_2026-06-30.xlsx").active
    assert pl["A3"].value == "1 July 2025 to 30 June 2026"
    assert [c.value for c in pl["A"]].count("Less Operating Expenses") == 1   # verbatim


def test_the_record_matches_the_bytes_on_disk(two, run_cli, tmp_path):
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 0
    record = json.loads((out / "PULL.json").read_text())
    assert record["as_at"] == AS_AT and record["version"] == env["version"]
    for org in record["organisations"]:
        assert org["authorised_by"] == "adviser@example.test"
        assert org["financial_year"] == {"start": "2025-07-01", "end": "2026-06-30"}
        for f in org["files"]:
            assert hashlib.sha256((out / f["file"]).read_bytes()).hexdigest() == f["sha256"]
        for raw in org["raw"]:
            assert hashlib.sha256((out / raw["file"]).read_bytes()).hexdigest() == raw["sha256"]
        assert org["calls"] == len(two.api_calls(org["tenant_id"]))
    text = (out / "PULL.md").read_text()
    assert "Account Transactions" in text and "ageing columns" in text
    assert "## A — Entity A Pty Ltd" in text


def test_raw_responses_are_kept_exactly_as_sent(two, run_cli, tmp_path):
    out = tmp_path / "p"
    _pull(run_cli, out)
    kept = (out / "raw" / "ACME_A_TrialBalance_2026-06-30.json").read_bytes()
    assert kept == two._trial_balance(two.orgs["t-a"], AS_AT)


def test_every_call_names_its_organisation_and_asks_for_json(two, run_cli, tmp_path):
    _pull(run_cli, tmp_path / "p")
    calls = two.api_calls()
    assert calls
    for method, url, headers in calls:
        assert method == "GET"
        assert headers["Xero-Tenant-Id"] in ("t-a", "t-b"), url
        assert headers["Accept"] == "application/json"
    assert "standardLayout=true" in next(u for _, u, _ in calls if "BalanceSheet" in u)


def test_compare_sets_two_dates_side_by_side(xero, run_cli, tmp_path):
    org = xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    # a month earlier: a term deposit that has since gone, and no receivable yet
    org.at["2026-05-31"] = Ledger(
        bank=[("090", "Business Bank Account", 3000), ("091", "Term Deposit", 7000)],
        current_assets=[], current_liabilities=[("800", "Accounts Payable", 3000)],
        income=[("200", "Sales", 30000)],
        expenses=[("400", "Advertising", 6000), ("477", "Wages and Salaries", 21000)])
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    out = tmp_path / "p"
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--compare", "2026-05-31",
                         "--org", "A", "--reports", "bs", "--out", str(out)])
    assert code == 0, env["problems"]
    name = "A_Balance_Sheet_as_at_2026-06-30_vs_2026-05-31.xlsx"
    ws = openpyxl.load_workbook(out / name).active
    assert [c.value for c in ws[5]] == ["Account", "30 Jun 2026", "31 May 2026"]
    row = {c.value: i for i, c in enumerate(ws["A"], start=1)}
    assert ws.cell(row=row["Business Bank Account"], column=2).value == 5000
    assert ws.cell(row=row["Business Bank Account"], column=3).value == 3000
    # present at one date only: the other cell is empty, never zero
    assert ws.cell(row=row["Term Deposit"], column=2).value is None
    assert ws.cell(row=row["Term Deposit"], column=3).value == 7000
    assert ws.cell(row=row["Accounts Receivable"], column=3).value is None
    assert row["Term Deposit"] < row["Total Bank"], "a line belongs above its total"
    code, env = run_cli(["read-ledger", str(out / name), "--expect-date", AS_AT])
    assert code == 0


def test_compare_must_be_earlier(two, run_cli, tmp_path):
    code, env = _pull(run_cli, tmp_path / "p", "--compare", "2026-07-31")
    assert code == 1 and env["problems"][0]["code"] == "COMPARE_NOT_EARLIER"


def test_the_financial_year_is_the_organisations_own(xero, run_cli, tmp_path):
    xero.add_org(Org("t-d", "Entity D Pty Ltd", fy_end=(31, 12)))
    signed_in(xero, ["t-d"], keys={"t-d": "D"})
    out = tmp_path / "p"
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "D", "--reports", "pl",
                         "--out", str(out)])
    assert code == 0, env["problems"]
    assert "D_Profit_and_Loss_2026-01-01_to_2026-06-30.xlsx" in _files(out)
    assert "D_Profit_and_Loss_2025-01-01_to_2025-12-31.xlsx" in _files(out)


# -- anything wrong, nothing written --------------------------------------------------

def test_a_title_that_echoes_another_date_writes_nothing(two, run_cli, tmp_path):
    two.skew_title[("t-b", "bs")] = "As at 31 July 2026"      # the end-of-month default
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "DATE_MISMATCH"
    assert "31 July 2026" in env["problems"][0]["message"]
    assert not out.exists()


def test_a_period_whose_start_is_wrong_writes_nothing(two, run_cli, tmp_path):
    two.skew_title[("t-a", "pl")] = "1 May 2026 to 30 June 2026"
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "DATE_MISMATCH"
    assert not out.exists()


def test_a_title_with_no_date_cannot_be_proved(two, run_cli, tmp_path):
    two.skew_title[("t-a", "tb")] = "Current"
    code, env = _pull(run_cli, tmp_path / "p")
    assert code == 1 and env["problems"][0]["code"] == "DATE_NOT_ECHOED"


def test_a_title_naming_the_end_and_the_right_span_proves_the_period(two, run_cli, tmp_path):
    two.skew_title[("t-a", "bank")] = "For the month ended 30 June 2026"
    code, env = _pull(run_cli, tmp_path / "p")
    assert code == 0
    assert not any("only the end of the period" in w for w in env["warnings"])


def test_a_title_naming_only_an_end_is_accepted_and_said(two, run_cli, tmp_path):
    two.skew_title[("t-a", "bank")] = "Period ending 30 June 2026"
    code, env = _pull(run_cli, tmp_path / "p")
    assert code == 0
    assert any("A Bank Summary" in w and "only the end of the period" in w
               for w in env["warnings"])


def test_a_year_to_date_answered_with_a_month_is_refused(two, run_cli, tmp_path):
    # the end matches; the span the words name does not
    two.skew_title[("t-a", "pl")] = "For the month ended 30 June 2026"
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "DATE_MISMATCH"
    assert "a period starting 2026-06-01" in env["problems"][0]["message"]
    assert not out.exists()


def test_a_company_whose_name_reads_as_a_date_is_refused_before_any_report(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    xero.add_org(Org("t-o", "Old File to 30 June 2023"))
    signed_in(xero, ["t-a", "t-o"], keys={"t-a": "A", "t-o": "OLD"})
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "ORG_NAME_READS_AS_A_DATE"
    assert "OLD" in env["problems"][0]["message"]
    assert not out.exists()
    assert all("Organisation" in u for _, u, _ in xero.api_calls()), \
        "refused at the first look, before a report was asked for"
    # the others can still be pulled by name
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "A", "--reports", "tb",
                         "--out", str(out)])
    assert code == 0


def test_a_name_with_a_stray_space_is_not_a_reason_to_refuse(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd "))
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "A",
                         "--out", str(tmp_path / "p")])
    assert code == 0, env["problems"]


def test_a_folder_that_cannot_be_written_is_exit_2_before_xero_is_asked(two, run_cli, tmp_path):
    import os
    locked = tmp_path / "locked"
    locked.mkdir()
    os.chmod(locked, 0o500)
    try:
        code, env = _pull(run_cli, locked / "p")
    finally:
        os.chmod(locked, 0o700)
    assert code == 2 and env["problems"][0]["code"] == "CANNOT_WRITE"
    assert not two.api_calls()


def test_a_total_that_does_not_foot_writes_nothing_and_names_every_break(two, run_cli, tmp_path):
    two.bend_total[("t-a", "bs", "Total Bank")] = 10
    two.bend_total[("t-b", "pl", "Total Income")] = -250
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["status"] == "refuse"
    messages = " | ".join(p["message"] for p in env["problems"])
    assert "A Balance Sheet: Total Bank" in messages and "apart by -10.00" in messages
    assert "B Profit and Loss" in messages and "Total Income" in messages
    assert not out.exists(), "one break anywhere and the folder is left as it was"


def test_profit_that_disagrees_with_current_year_earnings_refuses(two, run_cli, tmp_path):
    two.bend_total[("t-a", "pl", "Net Profit")] = 500
    code, env = _pull(run_cli, tmp_path / "p")
    assert code == 1
    assert any("year-to-date Net Profit = Current Year Earnings" in p["message"]
               for p in env["problems"])


def test_an_out_of_balance_trial_balance_refuses(two, run_cli, tmp_path):
    two.one_sided["t-a"] = 77                          # a debit with no credit
    out = tmp_path / "p"
    code, env = _pull(run_cli, out, "--reports", "tb")
    assert code == 1
    messages = " | ".join(p["message"] for p in env["problems"])
    assert "A Trial Balance: Debit = Credit" in messages and "apart by 77.00" in messages
    assert "YTD Debit = YTD Credit" in messages
    assert not out.exists()


def test_xeros_own_total_line_must_be_the_sum_of_the_accounts(two, run_cli, tmp_path):
    two.bend_total[("t-a", "tb", "Total")] = 99
    code, env = _pull(run_cli, tmp_path / "p", "--reports", "tb")
    assert code == 1
    assert any("A Trial Balance: Total (Debit)" in p["message"] for p in env["problems"])


def test_a_cent_of_rounding_passes_and_is_said(two, run_cli, tmp_path):
    two.bend_total[("t-a", "bs", "Total Bank")] = Decimal("0.01")
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 0
    assert any("within rounding" in w and "Total Bank" in w for w in env["warnings"])
    record = json.loads((out / "PULL.json").read_text())
    assert any(t["status"] == "rounding" for t in record["organisations"][0]["ties"])


def test_an_allowance_too_small_refuses_before_anything_is_written(two, run_cli, tmp_path):
    two.day_remaining = 4
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "XERO_DAILY_LIMIT"
    assert "needs 7 more calls" in env["problems"][0]["message"]
    assert not out.exists()


def test_a_missing_permission_names_itself(two, run_cli, tmp_path):
    two.no_scope_for.add("Reports/BankSummary")
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 1 and env["problems"][0]["code"] == "XERO_SCOPE_MISSING"
    assert "Reports/BankSummary" in env["problems"][0]["message"]
    assert not out.exists()


def test_a_failed_write_removes_everything_this_run_made(two, run_cli, tmp_path, monkeypatch):
    from fma_tools.xero import layout
    real, seen = layout.write_workbook, []

    def third_one_fails(path, sheets, expect=None):
        seen.append(path)
        if len(seen) == 3:
            raise layout.InputProblem("CANNOT_WRITE", "disk went away (forced)")
        return real(path, sheets, expect)
    monkeypatch.setattr(layout, "write_workbook", third_one_fails)
    out = tmp_path / "p"
    code, env = _pull(run_cli, out)
    assert code == 2 and env["problems"][0]["code"] == "CANNOT_WRITE"
    assert not out.exists()


def test_a_label_that_starts_with_equals_stays_text(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd", ledger=Ledger(
        expenses=[("400", "=Advertising", 6000), ("404", "Bank Fees", 1000),
                  ("477", "Wages and Salaries", 21000)])))
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    out = tmp_path / "p"
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--org", "A", "--reports", "pl",
                         "--out", str(out)])
    assert code == 0, env["problems"]
    path = out / "A_Profit_and_Loss_2025-07-01_to_2026-06-30.xlsx"
    code, env = run_cli(["read-ledger", str(path), "--expect-date", AS_AT])
    assert code == 0 and env["data"]["sheets"][0]["formula_count"] == 0
    assert any("=Advertising" in row for row in env["data"]["sheets"][0]["rows"])


def test_no_token_reaches_any_output(two, run_cli, tmp_path, capsys):
    from fma_tools.cli import main
    out = tmp_path / "p"
    assert main(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(out)]) == 0
    captured = capsys.readouterr()
    blob = captured.out + captured.err
    for p in out.rglob("*"):
        if p.is_file() and p.suffix in (".json", ".md"):
            blob += p.read_text()
    assert two.issued
    for token in two.issued:
        assert token not in blob


# -- the pieces ----------------------------------------------------------------------

@pytest.mark.parametrize("as_at, month, day, want", [
    ("2026-06-30", 6, 30, ("2025-07-01", "2026-06-30")),
    ("2026-07-01", 6, 30, ("2026-07-01", "2027-06-30")),
    ("2026-09-30", 12, 31, ("2026-01-01", "2026-12-31")),
    ("2024-02-29", 2, 29, ("2023-03-01", "2024-02-29")),
    ("2025-02-28", 2, 29, ("2024-03-01", "2025-02-28")),
    # a February year end stored as the 28th still takes in the 29th of a leap year
    ("2028-02-29", 2, 28, ("2027-03-01", "2028-02-29")),
    ("2028-03-01", 2, 28, ("2028-03-01", "2029-02-28")),
    ("2027-02-28", 2, 28, ("2026-03-01", "2027-02-28")),
    ("2026-03-15", 3, 31, ("2025-04-01", "2026-03-31")),
])
def test_financial_year_bounds(as_at, month, day, want):
    start, end = reports.financial_year(date.fromisoformat(as_at), month, day)
    assert (start.isoformat(), end.isoformat()) == want


@pytest.mark.parametrize("line, want", [
    ("As at 30 June 2026", ["2026-06-30"]),
    ("1 July 2025 to 30 June 2026", ["2025-07-01", "2026-06-30"]),
    ("From 1 Jun 2026 to 30 Jun 2026", ["2026-06-01", "2026-06-30"]),
    ("For the month ended 30 September 2026", ["2026-09-30"]),
    ("Entity A Pty Ltd", []),
    ("31 February 2026", []),
])
def test_dates_are_read_from_a_title_line(line, want):
    assert [d.isoformat() for d in reports.dates_in(line)] == want


@pytest.mark.parametrize("text, want", [
    ("1234.50", Decimal("1234.50")), ("-0.01", Decimal("-0.01")),
    ("1,234.50", Decimal("1234.50")), ("(250.00)", Decimal("-250.00")),
    ("", None), (None, None), ("n/a", "n/a"), ("NaN", "NaN"),
])
def test_cell_text_becomes_decimal_or_stays_what_it_was(text, want):
    assert reports.to_number(text) == want
    assert not isinstance(reports.to_number(text), float)


def test_long_dates_do_not_depend_on_the_locale():
    assert reports.as_at_line(date(2026, 9, 3)) == "As at 3 September 2026"
    assert reports.range_line(date(2025, 7, 1), date(2026, 6, 30)) == \
        "1 July 2025 to 30 June 2026"


def test_three_companies_pull_in_one_go(xero, run_cli, tmp_path):
    three_orgs(xero)
    signed_in(xero, ["t-a", "t-b", "t-c"], keys={"t-a": "A", "t-b": "B", "t-c": "C"})
    code, env = _pull(run_cli, tmp_path / "p")
    assert code == 0, env["problems"]
    assert [o["key"] for o in env["data"]["organisations"]] == ["A", "B", "C"]
    assert all(len(o["files"]) == 7 for o in env["data"]["organisations"])


def test_in_the_first_month_of_the_year_the_earnings_tie_still_runs(two, run_cli, tmp_path):
    # 1-31 July is both the month to date and the year to date for a June year end:
    # fetched once, and still tied to Current Year Earnings.
    out = tmp_path / "p"
    code, env = run_cli(["xero", "pull", "--as-at", "2026-07-31", "--org", "A",
                         "--out", str(out)])
    assert code == 0, env["problems"]
    names = _files(out)
    assert "A_Profit_and_Loss_2026-07-01_to_2026-07-31.xlsx" in names
    assert sum("Profit_and_Loss" in n for n in names) == 2        # that one, and the prior year
    record = json.loads((out / "PULL.json").read_text())
    ties = [t for t in record["organisations"][0]["ties"]
            if "Current Year Earnings" in t["name"]]
    assert len(ties) == 1 and ties[0]["status"] == "ok"


def test_compare_without_a_balance_sheet_is_said(two, run_cli, tmp_path):
    code, env = _pull(run_cli, tmp_path / "p", "--compare", "2026-05-31", "--reports", "tb")
    assert code == 0
    assert any("leaves the balance sheet out" in w for w in env["warnings"])
