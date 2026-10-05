"""xero group: side by side, summed, and nothing claimed that was not done."""

import json

import openpyxl
import pytest

from xero_fakes import Ledger, Org, signed_in, three_orgs

AS_AT = "2026-06-30"


@pytest.fixture
def pulled(xero, run_cli, tmp_path):
    """Three invented companies, pulled. Returns (pull folder, fake)."""
    three_orgs(xero)
    signed_in(xero, ["t-a", "t-b", "t-c"], keys={"t-a": "A", "t-b": "B", "t-c": "C"})
    out = tmp_path / "pull"
    code, env = run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--prefix", "ACME",
                         "--out", str(out)])
    assert code == 0, env["problems"]
    return out


def _group(run_cli, pull, out, *extra, as_at=AS_AT):
    return run_cli(["xero", "group", "--pull", str(pull), "--as-at", as_at,
                    "--out", str(out), *extra])


def _table(ws) -> tuple[list, dict]:
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    header = rows[4]
    by_code = {r[0]: dict(zip(header, r)) for r in rows[5:] if r[0]}
    return rows, by_code


def test_the_sheet_says_what_it_is_and_what_it_is_not(pulled, run_cli, tmp_path):
    out = tmp_path / "group.xlsx"
    code, env = _group(run_cli, pulled, out, "--group-name", "ACME group")
    assert code == 0, env["problems"]
    ws = openpyxl.load_workbook(out)["Group"]
    assert ws["A1"].value == "Group trial balance"
    assert ws["A2"].value == ("ACME group: management aggregation of 3 organisations, "
                              "no eliminations, not statutory accounts")
    assert ws["A3"].value == "As at 30 June 2026"
    assert env["data"]["eliminations"] == "none; the column is empty"


def test_each_company_keeps_its_column_and_the_group_is_their_sum(pulled, run_cli, tmp_path):
    out = tmp_path / "group.xlsx"
    code, env = _group(run_cli, pulled, out)
    assert code == 0
    rows, by_code = _table(openpyxl.load_workbook(out)["Group"])
    header = rows[4]
    assert header == ["Code", "Account", "Class", "Type", "Entity A Pty Ltd",
                      "Entity B Pty Ltd", "Entity C Pty Ltd", "Group (sum)",
                      "Eliminations", "Intercompany"]
    bank = by_code["090"]
    assert (bank["Entity A Pty Ltd"], bank["Entity B Pty Ltd"], bank["Entity C Pty Ltd"]) == \
        (5000, 2000, 1000)
    assert bank["Group (sum)"] == 8000
    assert by_code["200"]["Group (sum)"] == -52000        # revenue is a credit
    assert all(r["Eliminations"] is None for r in by_code.values())
    # an account one company does not have is an empty cell, not a zero
    assert by_code["405"]["Entity A Pty Ltd"] is None
    assert by_code["405"]["Entity B Pty Ltd"] == 500
    total = next(r for r in rows if r[1] == "Total (debits less credits)")
    assert total[4:8] == [0, 0, 0, 0], "every trial balance nets to nothing, and so does the sum"
    # classes in statement order
    classes = [r["Class"] for r in by_code.values()]
    assert classes == sorted(classes, key=["ASSET", "LIABILITY", "EQUITY", "REVENUE",
                                           "EXPENSE"].index)


def test_chart_differences_are_listed_never_guessed(pulled, run_cli, tmp_path):
    out = tmp_path / "group.xlsx"
    code, env = _group(run_cli, pulled, out)
    assert code == 0 and env["data"]["differences"] >= 4
    rows = [[c.value for c in r] for r in openpyxl.load_workbook(out)["Chart differences"].iter_rows()]
    found = {(r[0], r[1]): r[3] for r in rows[5:]}
    assert found[("In some organisations only", "405")] == "in B; not in A, C"
    assert found[("In some organisations only", "404")] == "in A; not in B, C"
    assert "A: Advertising; B: Advertising; C: Marketing" == \
        found[("One code, different names", "400")]
    assert ("One name, different codes", "404, 410") in found      # "Bank Fees"
    # and the Group sheet did NOT quietly merge 404 with 410
    _, by_code = _table(openpyxl.load_workbook(out)["Group"])
    assert by_code["404"]["Entity C Pty Ltd"] is None and by_code["410"]["Entity C Pty Ltd"] == 100


def test_a_mapping_lines_codes_up_and_flags_intercompany(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code,group_name,intercompany\n"
                       "C,410,404,Bank Fees,\n"
                       "b,405,404,Bank Fees,no\n"
                       "A,610,610,Receivables,yes\n")
    out = tmp_path / "group.xlsx"
    code, env = _group(run_cli, pulled, out, "--mapping", str(mapping))
    assert code == 0, env["problems"]
    assert env["data"]["mapping_lines"] == 3
    assert env["data"]["mapping_lines_with_a_balance"] == 3
    _, by_code = _table(openpyxl.load_workbook(out)["Group"])
    fees = by_code["404"]
    assert (fees["Entity A Pty Ltd"], fees["Entity B Pty Ltd"], fees["Entity C Pty Ltd"]) == \
        (1000, 500, 100)
    assert fees["Group (sum)"] == 1600 and "410" not in by_code and "405" not in by_code
    assert by_code["610"]["Intercompany"] == "yes" and fees["Intercompany"] is None


def test_a_mapping_that_names_an_unknown_company_refuses(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code\nZ,410,404\n")
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 1 and env["problems"][0]["code"] == "MAPPING_UNKNOWN_ENTITY"
    mapping.write_text("company,account\nA,1\n")
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 2 and env["problems"][0]["code"] == "MAPPING_INVALID"


def test_the_group_sheet_passes_the_same_date_gate(pulled, run_cli, tmp_path):
    out = tmp_path / "group.xlsx"
    assert _group(run_cli, pulled, out)[0] == 0
    code, env = run_cli(["read-ledger", str(out), "--expect-date", AS_AT])
    assert code == 0, env["problems"]
    assert env["data"]["metadata"]["report_title"] == "Group trial balance"
    # it is told apart from a pull: a derived sheet, never "Xero's own figures"
    assert "written_by" not in env["data"]["metadata"]
    assert len(env["warnings"]) == 1 and "not an export from Xero" in env["warnings"][0]
    code, env = run_cli(["read-ledger", str(out), "--expect-date", "2026-05-31"])
    assert code == 1


def test_the_date_is_checked_against_the_pull(pulled, run_cli, tmp_path):
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", as_at="2026-05-31")
    assert code == 1 and env["problems"][0]["code"] == "DATE_MISMATCH"
    code, env = run_cli(["xero", "group", "--pull", str(pulled), "--out",
                         str(tmp_path / "g.xlsx")])
    assert code == 1 and env["problems"][0]["code"] == "AS_AT_REQUIRED"


def test_a_folder_that_is_not_a_pull_refuses(run_cli, tmp_path):
    code, env = _group(run_cli, tmp_path, tmp_path / "g.xlsx")
    assert code == 1 and env["problems"][0]["code"] == "PULL_RECORD_MISSING"


def test_inputs_changed_since_the_pull_refuse(pulled, run_cli, tmp_path):
    raw = pulled / "raw" / "ACME_B_TrialBalance_2026-06-30.json"
    raw.write_bytes(raw.read_bytes().replace(b"2000.00", b"2999.00"))
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx")
    assert code == 1 and env["problems"][0]["code"] == "PULL_CHANGED"
    assert not (tmp_path / "g.xlsx").exists()


def test_an_out_of_balance_company_is_not_added_to_anything(pulled, run_cli, tmp_path):
    record_path = pulled / "PULL.json"
    doc = json.loads(record_path.read_text())
    raw = pulled / "raw" / "ACME_B_TrialBalance_2026-06-30.json"
    blob = raw.read_bytes().replace(b'"Value": "500.00"', b'"Value": "590.00"')
    raw.write_bytes(blob)
    # make the record agree with the changed bytes, so only the imbalance is at issue
    import hashlib
    for org in doc["organisations"]:
        for r in org["raw"]:
            if r["file"].endswith("ACME_B_TrialBalance_2026-06-30.json"):
                r["sha256"] = hashlib.sha256(blob).hexdigest()
    record_path.write_text(json.dumps(doc))
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx")
    assert code == 1 and env["problems"][0]["code"] == "TRIAL_BALANCE_OUT"
    assert "B:" in env["problems"][0]["message"]


def test_different_currencies_cannot_be_added(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    xero.add_org(Org("t-n", "Entity N Limited", currency="NZD"))
    signed_in(xero, ["t-a", "t-n"], keys={"t-a": "A", "t-n": "N"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(pull)])[0] == 0
    code, env = _group(run_cli, pull, tmp_path / "g.xlsx")
    assert code == 1 and env["problems"][0]["code"] == "CURRENCIES_DIFFER"
    assert "AUD, NZD" in env["problems"][0]["message"]


def test_different_year_ends_are_said_on_the_face(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    xero.add_org(Org("t-d", "Entity D Pty Ltd", fy_end=(31, 12)))
    signed_in(xero, ["t-a", "t-d"], keys={"t-a": "A", "t-d": "D"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(pull)])[0] == 0
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pull, out)
    assert code == 0
    assert any("different financial-year ends" in w for w in env["warnings"])
    assert "financial-year ends differ" in openpyxl.load_workbook(out)["Group"]["A2"].value


@pytest.mark.parametrize("kinds, missing", [("bs,pl,coa", "trial balance"),
                                            ("tb", "chart of accounts")])
def test_a_group_sheet_needs_the_trial_balance_and_the_chart(xero, run_cli, tmp_path,
                                                             kinds, missing):
    # Without the chart, an account's code could only be guessed from its name.
    xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--reports", kinds,
                    "--out", str(pull)])[0] == 0
    code, env = _group(run_cli, pull, tmp_path / "g.xlsx")
    assert code == 1 and env["problems"][0]["code"] == "PULL_INCOMPLETE"
    assert f"no {missing} for A" in env["problems"][0]["message"]
    assert not (tmp_path / "g.xlsx").exists()


def test_group_uses_no_network(pulled, run_cli, tmp_path, monkeypatch):
    from fma_tools.xero import transport

    def none():
        raise AssertionError("group must not touch the network")
    monkeypatch.setattr(transport, "_FACTORY", none)
    assert _group(run_cli, pulled, tmp_path / "g.xlsx")[0] == 0


def test_one_company_is_still_a_sheet(xero, run_cli, tmp_path):
    xero.add_org(Org("t-a", "Entity A Pty Ltd", ledger=Ledger()))
    signed_in(xero, ["t-a"], keys={"t-a": "A"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(pull)])[0] == 0
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pull, out)
    assert code == 0 and env["data"]["differences"] == 0
    about = [[c.value for c in r] for r in openpyxl.load_workbook(out)["About"].iter_rows()]
    assert about[4][:3] == ["Key", "Legal entity", "Xero name"]
    assert about[5][0] == "A" and about[5][5] == "adviser@example.test"


def test_accounts_with_no_code_are_never_lined_up_by_what_their_names_contain(xero, run_cli, tmp_path):
    # Bank accounts in Xero often have no code. "(USD)" and "(800)" in a NAME are not
    # codes: two such accounts are not one account, and neither is Accounts Payable.
    xero.add_org(Org("t-a", "Entity A Pty Ltd", ledger=Ledger(
        bank=[("", "Wise (USD)", 5000), ("", "Term Deposit (800)", 2000)])))
    xero.add_org(Org("t-b", "Entity B Pty Ltd", ledger=Ledger(
        bank=[("", "PayPal (USD)", 2000), ("", "Wise (USD)", 300)],
        current_assets=[("610", "Accounts Receivable", 4000)],
        current_liabilities=[("800", "Accounts Payable", 1000)],
        income=[("200", "Sales", 9000)], expenses=[("400", "Advertising", 2000)])))
    signed_in(xero, ["t-a", "t-b"], keys={"t-a": "A", "t-b": "B"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(pull)])[0] == 0
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pull, out)
    assert code == 0, env["problems"]
    wb = openpyxl.load_workbook(out)
    rows = [[c.value for c in r] for r in wb["Group"].iter_rows()]
    header = rows[4]
    a, b = header.index("Entity A Pty Ltd"), header.index("Entity B Pty Ltd")
    lines = [r for r in rows[5:] if r[1] and r[1] != "Total (debits less credits)"]
    by_name = {}
    for r in lines:
        by_name.setdefault(r[1], []).append(r)
    # no row was given a code out of a name
    assert not [r for r in lines if r[0] in ("USD", "800") and r[1] != "Accounts Payable"]
    # each code-less account is its own line, in its own company's column only
    assert [(r[0], r[a], r[b]) for r in by_name["Term Deposit (800)"]] == [(None, 2000, None)]
    assert [(r[0], r[a], r[b]) for r in by_name["PayPal (USD)"]] == [(None, None, 2000)]
    assert {(r[a], r[b]) for r in by_name["Wise (USD)"]} == {(5000, None), (None, 300)}, \
        "same name, two companies, two lines"
    assert len(by_name["Wise (USD)"]) == 2
    payable = by_name["Accounts Payable"][0]
    assert payable[0] == "800" and (payable[a], payable[b]) == (-3000, -1000)
    assert payable[2] == "LIABILITY", "a bank asset was not netted into it"
    # and a person is told
    diffs = [[c.value for c in r] for r in wb["Chart differences"].iter_rows()][5:]
    said = [d for d in diffs if d[0] == "Not lined up"]
    assert sorted(d[2] for d in said) == ["PayPal (USD)", "Term Deposit (800)",
                                          "Wise (USD)", "Wise (USD)"]
    assert all("the account has no code" in d[3] for d in said)
    total = next(r for r in rows if r[1] == "Total (debits less credits)")
    assert total[a] == 0 and total[b] == 0


def test_a_sheet_that_exists_is_never_written_over_without_being_told(pulled, run_cli, tmp_path):
    out = tmp_path / "g.xlsx"
    assert _group(run_cli, pulled, out)[0] == 0
    wb = openpyxl.load_workbook(out)                  # someone types an elimination in
    ws = wb["Group"]
    header = [c.value for c in ws[5]]
    ws.cell(row=6, column=header.index("Eliminations") + 1, value=-123.45)
    wb.save(out)
    before = out.read_bytes()
    code, env = _group(run_cli, pulled, out)
    assert code == 1 and env["problems"][0]["code"] == "OUT_EXISTS"
    assert out.read_bytes() == before
    code, env = _group(run_cli, pulled, out, "--replace")
    assert code == 0 and out.read_bytes() != before


def test_a_mapping_line_that_matches_no_account_refuses(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code\nC,41O,404\nA,404,404\n")   # letter O
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pulled, out, "--mapping", str(mapping))
    assert code == 1 and env["problems"][0]["code"] == "MAPPING_UNMATCHED"
    assert env["problems"][0]["message"] == "line 2: C has no account coded '41O'"
    assert not out.exists()


def test_a_mapping_that_excel_saved_badly_is_exit_2_not_a_bug(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_bytes("entity,code,group_code,group_name\nA,404,404,Café fees\n"
                        .encode("cp1252"))
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 2 and env["problems"][0]["code"] == "MAPPING_INVALID"
    assert "CSV UTF-8" in env["problems"][0]["message"]
    mapping.write_text("entity,code,group_code\nA,404,404,one,cell,too,many\n")
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 2 and "line 2 has more cells than the header" in env["problems"][0]["message"]


def test_an_archived_account_that_still_holds_a_balance_is_accounted_for(xero, run_cli, tmp_path):
    # A's 470 is archived but carries 900; B's 470 is a different, active account. The
    # balances share the line (the code is the rule) -- so the sheet that exists to
    # flag exactly this must say so, not claim A has no 470.
    a = xero.add_org(Org("t-a", "Entity A Pty Ltd"))
    a.ledger.expenses.append(("470", "Old Motor Vehicle Costs", 900))
    xero.add_org(Org("t-b", "Entity B Pty Ltd", ledger=Ledger(
        expenses=[("400", "Advertising", 6000), ("404", "Bank Fees", 1000),
                  ("470", "Subscriptions", 500), ("477", "Wages and Salaries", 21000)])))
    a.archive_codes = {"470"}
    signed_in(xero, ["t-a", "t-b"], keys={"t-a": "A", "t-b": "B"})
    pull = tmp_path / "pull"
    assert run_cli(["xero", "pull", "--as-at", AS_AT, "--all", "--out", str(pull)])[0] == 0
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pull, out)
    assert code == 0, env["problems"]
    wb = openpyxl.load_workbook(out)
    _, by_code = _table(wb["Group"])
    assert (by_code["470"]["Entity A Pty Ltd"], by_code["470"]["Entity B Pty Ltd"]) == (900, 500)
    diffs = {(r[0].value, r[1].value): r[3].value for r in wb["Chart differences"].iter_rows(min_row=6)}
    assert ("One code, different names", "470") in diffs
    assert "A: Old Motor Vehicle Costs; B: Subscriptions" == diffs[("One code, different names", "470")]
    assert ("In some organisations only", "470") not in diffs, "A does have a 470"


def test_a_group_code_that_looks_like_a_typo_is_said(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code\nC,410,4O4\nB,405,404\n")     # letter O
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pulled, out, "--mapping", str(mapping))
    assert code == 0, "a group code may be new on purpose, so this cannot refuse"
    assert any("mapping line 2 sends C 410 to group code '4O4'" in w and "typo" in w
               for w in env["warnings"])
    assert not any("line 3" in w for w in env["warnings"]), "404 is a real code"


def test_mapping_line_numbers_are_the_ones_a_person_sees(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code,group_name\n"
                       "A,404,404,Bank fees\n"
                       "\n"
                       "\n"
                       "C,41O,404,\n")
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 1
    assert env["problems"][0]["message"] == "line 5: C has no account coded '41O'"


def test_a_mapping_with_a_header_and_nothing_else_maps_nothing(pulled, run_cli, tmp_path):
    mapping = tmp_path / "mapping.csv"
    mapping.write_text("entity,code,group_code\n")
    code, env = _group(run_cli, pulled, tmp_path / "g.xlsx", "--mapping", str(mapping))
    assert code == 0 and env["data"]["mapping_lines"] == 0


def test_an_out_that_is_a_folder_is_exit_2(pulled, run_cli, tmp_path):
    folder = tmp_path / "group.xlsx"
    folder.mkdir()
    code, env = _group(run_cli, pulled, folder, "--replace")
    assert code == 2 and env["problems"][0]["code"] == "OUT_IS_A_FOLDER"


def test_an_account_missing_from_the_chart_is_explained_as_that(pulled, run_cli, tmp_path):
    import hashlib
    record_path = pulled / "PULL.json"
    doc = json.loads(record_path.read_text())
    raw = pulled / "raw" / "ACME_B_Accounts.json"
    chart = json.loads(raw.read_text())
    chart["Accounts"] = [a for a in chart["Accounts"] if a["Code"] != "610"]
    blob = json.dumps(chart).encode()
    raw.write_bytes(blob)
    for org in doc["organisations"]:
        for r in org["raw"]:
            if r["file"].endswith("ACME_B_Accounts.json"):
                r["sha256"] = hashlib.sha256(blob).hexdigest()
    record_path.write_text(json.dumps(doc))
    out = tmp_path / "g.xlsx"
    code, env = _group(run_cli, pulled, out)
    assert code == 0
    diffs = [[c.value for c in r] for r in
             openpyxl.load_workbook(out)["Chart differences"].iter_rows(min_row=6)]
    alone = [d for d in diffs if d[0] == "Not lined up"]
    assert len(alone) == 1 and "not in the chart that was pulled" in alone[0][3]
    assert "in B" in alone[0][3]
